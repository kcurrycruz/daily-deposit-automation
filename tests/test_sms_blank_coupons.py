"""Blank coupon allocations must retain the pre-audit SMS workflow."""

from dataclasses import replace
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from app import pos_to_quickbooks_v2 as engine
from app.daily_reporting_workbook import build_reporting_workbook
from app.sms_deposit_data import build_sms_deposit_data
from app.sms_exports import build_sms_export_bundle
from tests.sms_fixture_factory import sms_exports_091426


def coupon_bundle(blank="", control="12.50"):
    # Legacy SMS detail Amount is one column left of its printed header.
    # Quantities deliberately differ from monetary amounts.
    coupon_rows = (
        ("", "", "Sub-Department", "", "", "", "", "Qty", "Amount", "Pkg/Weight"),
        ("", 110, "Grocery", "", "", "", 12, "12.50", "", ""),
        ("", 142, "Bulk Herbs", "", "", "", 135, blank, "", ""),
        ("", 210, "Deli", "", "", "", 15, blank, "", ""),
        ("", 320, "Meat", "", "", "", 7, blank, "", ""),
        ("", "", "", "", "Total", "", 169, "", control, ""),
    )
    return build_sms_export_bundle(
        replace(report, rows=coupon_rows) if report.role == "coupons" else report
        for report in sms_exports_091426()
    )


class BlankCouponTests(unittest.TestCase):
    def test_first_blank_legacy_coupon_row_does_not_hide_later_allocations(self):
        bundle = coupon_bundle()
        reports = dict(bundle.reports)
        coupon = reports["coupons"]
        reports["coupons"] = replace(coupon, rows=(
            coupon.rows[0], coupon.rows[2], coupon.rows[1], *coupon.rows[3:],
        ))
        data = build_sms_deposit_data(replace(bundle, reports=reports))
        self.assertEqual(data.coupons_by_subdept, {110: Decimal("12.50")})
        self.assertEqual(data.coupon_total, Decimal("12.50"))
        output = build_reporting_workbook(
            Path("assets/SubDept Single Total Report Template.xlsx"),
            replace(bundle, reports=reports), data,
        )
        workbook = load_workbook(BytesIO(output), data_only=True)
        try:
            source = workbook["SubDept Coupon (Local Discount)"]
            self.assertEqual(source["A1"].value, 142)
            self.assertEqual(source["F1"].value, 135)
            self.assertIsNone(source["G1"].value)
            self.assertEqual(source["A2"].value, 110)
            self.assertEqual(source["G2"].value, 12.50)
        finally:
            workbook.close()

    def test_first_blank_legacy_row_cannot_hide_later_invalid_money_with_zero_control(self):
        for amount in ("12.50", "bad amount", "NaN", "Infinity"):
            with self.subTest(amount=amount):
                bundle = coupon_bundle(control="0.00")
                reports = dict(bundle.reports)
                coupon = reports["coupons"]
                occupied = list(coupon.rows[1])
                occupied[7] = amount
                reports["coupons"] = replace(coupon, rows=(
                    coupon.rows[0], coupon.rows[2], tuple(occupied), *coupon.rows[3:],
                ))
                with self.assertRaises(ValueError):
                    build_sms_deposit_data(replace(bundle, reports=reports))

    def test_first_invalid_legacy_coupon_amount_cannot_be_treated_as_blank(self):
        for amount in ("bad amount", "NaN", "Infinity", True):
            with self.subTest(amount=amount):
                bundle = coupon_bundle(control="0.00")
                reports = dict(bundle.reports)
                coupon = reports["coupons"]
                occupied = list(coupon.rows[1])
                occupied[7] = amount
                reports["coupons"] = replace(coupon, rows=(
                    coupon.rows[0], tuple(occupied), *coupon.rows[2:],
                ))
                with self.assertRaises(ValueError):
                    build_sms_deposit_data(replace(bundle, reports=reports))

    def test_reconciled_blank_coupon_allocations_are_skipped_like_pre_audit(self):
        for blank in (None, "", "   "):
            with self.subTest(blank=blank):
                data = build_sms_deposit_data(coupon_bundle(blank))
                self.assertEqual(data.coupons_by_subdept, {110: Decimal("12.50")})
                self.assertEqual(data.coupon_total, Decimal("12.50"))
                self.assertEqual(data.store_coupons_raw, Decimal("0.00"))

    def test_blank_coupon_allocations_do_not_hide_control_mismatch(self):
        with self.assertRaisesRegex(ValueError, "detail total.*control total"):
            build_sms_deposit_data(coupon_bundle(control="13.50"))

    def test_blank_coupon_allocations_require_a_real_monetary_control(self):
        for control in (None, "", "bad amount", "NaN", "Infinity"):
            with self.subTest(control=control), self.assertRaises(ValueError):
                build_sms_deposit_data(coupon_bundle(control=control))

    def test_populated_invalid_coupon_amounts_are_not_silently_skipped(self):
        for invalid in ("bad amount", "NaN", "Infinity", True):
            with self.subTest(invalid=invalid), self.assertRaisesRegex(
                ValueError, "code 142.*valid monetary amount"
            ):
                build_sms_deposit_data(coupon_bundle(blank=invalid))

    def test_explicit_store_coupon_amount_is_not_inferred_from_blank(self):
        bundle = coupon_bundle()
        reports = dict(bundle.reports)
        reports["coupons"] = replace(reports["coupons"], rows=(
            ("Code", "Sub-Department", "Qty", "Amount"),
            (3542, "Store Coupons", 10, None),
            ("Coupon Distribution Total", "", "", 0),
            ("Store Coupons Total", "", "", 0),
        ))
        with self.assertRaisesRegex(ValueError, "code 3542.*valid monetary amount"):
            build_sms_deposit_data(replace(bundle, reports=reports))

    def test_blank_coupons_allow_reporting_and_keep_existing_unmapped_tba_path(self):
        bundle = coupon_bundle()
        reports = dict(bundle.reports)
        reports["sales"] = replace(reports["sales"], rows=(
            ("Code", "Sub-Department", "Qty", "Amount"),
            (110, "Grocery", 1, "100.00"),
            (777, "New Department", 1, "9.84"),
            ("Sales Total", "", "", "109.84"),
        ))
        bundle = replace(bundle, reports=reports)
        data = build_sms_deposit_data(bundle)
        output = build_reporting_workbook(
            Path("assets/SubDept Single Total Report Template.xlsx"), bundle, data
        )
        workbook = load_workbook(BytesIO(output), data_only=True)
        try:
            report = workbook["SubDept Sales Report"]
            for code in (142, 210, 320):
                row = next(row for row in report.iter_rows(values_only=True) if row[0] == code)
                self.assertEqual(row[3], 0)
            self.assertEqual(report["J3"].value, 109.84)
        finally:
            workbook.close()

        with TemporaryDirectory() as folder:
            root = Path(folder)
            path = root / "report.xlsx"
            path.write_bytes(output)
            with patch.object(engine.log, "disabled", True), \
                    patch.object(engine, "output_dir", root), \
                    patch.object(engine, "LOG_DIR", root), \
                    patch.object(engine, "SUMMARY_DIR", root), \
                    patch.dict(engine.CONFIG, pos_export_folder=str(root)):
                parsed = engine.parse_excel_report(path, bundle.deposit_date)
                self.assertEqual(parsed[1], [("Dept 777", 9.84)])
                iif = engine.generate_iif(
                    parsed[0], {}, {}, bundle.deposit_date,
                    owner_local_amt=12.50, misc_tba_lines=parsed[1],
                    excel_sales_total=109.84, hash_sales_total=0,
                ).read_text(encoding="utf-8")
                self.assertIn("4444 · TBA Purchases\t\t-9.84\tDept 777", iif)


if __name__ == "__main__":
    unittest.main()
