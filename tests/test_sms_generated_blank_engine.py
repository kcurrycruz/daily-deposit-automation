"""Optional blank SMS rows remain harmless through QuickBooks preparation."""

from dataclasses import replace
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook

from app import pos_to_quickbooks_v2 as engine
from app.daily_reporting_workbook import build_reporting_workbook
from app.sms_deposit_data import build_sms_deposit_data
from app.sms_exports import build_sms_export_bundle
from tests.sms_fixture_factory import sms_exports_091426


def _generated_workbook_with_blank_rows():
    reports = []
    for report in sms_exports_091426():
        if report.role == "discounts":
            rows = (
                ("Discounts by Shopper Level",),
                ("Date:", "09/14/2026"),
                (None, None, "Description", "Shopper Level", None, None, None, None, "Discount"),
                (None, None, "Owner", 2, None, None, None, None, "1200.00"),
                (None, None, "Senior", 3, None, None, None, None, "2781.21"),
                (None, None, "Unused", 4, None, None, None, None, None),
                (None, None, "Member Discounts", None, None, None, None, None, "3981.21"),
            )
            report = replace(report, rows=rows)
        elif report.role == "hash":
            rows = (
                ("Sub-department Single Total",),
                ("Date:", "09/14/2026", "to", "09/14/2026"),
                (None, None, "Sub-Department", None, None, None, None, "Qty", "Amount"),
                (None, 99, "Unused HASH", None, None, None, 1, None, None),
                (None, 23, "Refunded Discounts", None, None, None, 1, "4.89", None),
                (None, 32, "Pass Through Donations", None, None, None, 1, "6.00", None),
                (None, 34, "Paid In", None, None, None, 1, "12.00", None),
                ("HASH Detail Total", None, None, None, None, None, None, None, "22.89"),
            )
            report = replace(report, rows=rows)
        elif report.role == "bs":
            rows = (
                ("Store Balance Sheet",),
                ("Date:", "09/14/2026"),
                ("Code", "Description", None, None, "Amount"),
                (910, "Milk Bottle Return", None, None, "8.00"),
                (980, "Rev. Prepaid card/voucher", None, None, None),
            )
            report = replace(report, rows=rows)
        reports.append(report)
    bundle = build_sms_export_bundle(reports)
    data = build_sms_deposit_data(bundle)
    workbook = build_reporting_workbook(
        Path(__file__).resolve().parents[1]
        / "assets" / "SubDept Single Total Report Template.xlsx",
        bundle,
        data,
    )
    return bundle, workbook


class GeneratedBlankEngineTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.bundle, workbook = _generated_workbook_with_blank_rows()
        cls.folder = TemporaryDirectory()
        cls.path = Path(cls.folder.name) / "generated.xlsx"
        cls.path.write_bytes(workbook)

    @classmethod
    def tearDownClass(cls):
        cls.folder.cleanup()

    def test_blank_discount_level_does_not_block_engine(self):
        discounts, total = engine.parse_excel_discounts(
            self.path, self.bundle.deposit_date
        )
        self.assertEqual(total, 3981.21)
        self.assertEqual(sum(discounts.values()), 3981.21)

    def test_blank_hash_row_is_not_mistaken_for_quantity(self):
        details = engine.parse_hash_sheet_details(self.path, self.bundle.deposit_date)
        self.assertEqual(details["detail_total"], 22.89)
        self.assertEqual(details["misc_lines"], [])
        self.assertEqual(details["paid_in_total"], 12.00)

    def test_blank_optional_balance_sheet_code_does_not_block_engine(self):
        balance = engine.parse_bs_sheet(self.path, self.bundle.deposit_date)
        self.assertEqual(balance["milk_bottle_return"], 8.00)
        self.assertEqual(balance["prepaid_card"], 0.00)

    def test_normalized_totals_still_match_source_controls(self):
        data = build_sms_deposit_data(self.bundle)
        self.assertEqual(data.discount_total, Decimal("3981.21"))
        self.assertEqual(data.hash_paid_in, Decimal("12.00"))
        self.assertEqual(data.bs_data["prepaid_card"], Decimal("0.00"))

    def test_missing_hash_control_amount_cannot_use_quantity(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "091426 Hash"
        sheet.append(["Code", "Description", "Qty", "Amount"])
        sheet.append([23, "Refunded Discounts", 1, 4.89])
        sheet.append([None, "Total", 1, None])
        with TemporaryDirectory() as folder:
            path = Path(folder) / "missing-control.xlsx"
            workbook.save(path)
            loaded = load_workbook(path, read_only=True, data_only=True)
            try:
                with self.assertRaisesRegex(ValueError, "HASH control total"):
                    engine._hash_control_total(loaded, self.bundle.deposit_date)
            finally:
                loaded.close()


if __name__ == "__main__":
    unittest.main()
