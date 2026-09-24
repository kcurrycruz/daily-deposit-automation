import ast
import unittest
from dataclasses import replace
from datetime import date
from pathlib import Path
from uuid import uuid4

from app.daily_reporting_workbook import build_reporting_workbook
from app.sms_deposit_data import build_sms_deposit_data
from app.sms_exports import build_sms_export_bundle
from tests.sms_fixture_factory import sms_exports_091426


SOURCE = Path(__file__).resolve().parents[1] / "streamlit_app.py"


def streamlit_definition(name, namespace):
    node = next(
        node
        for node in ast.parse(SOURCE.read_text(encoding="utf-8")).body
        if isinstance(node, (ast.ClassDef, ast.FunctionDef)) and node.name == name
    )
    exec(
        compile(ast.Module(body=[node], type_ignores=[]), str(SOURCE), "exec"),
        namespace,
    )
    return namespace[name]


class DepositWorkflowTests(unittest.TestCase):
    def test_section_status_stops_before_the_next_validation_section(self):
        from typing import Optional

        section_status = streamlit_definition(
            "section_status",
            {"Optional": Optional},
        )
        log_text = """SALES CHECK
Script Net Sales: $88,463.13
Excel Sales Total: $88,463.13
RESULT: ✓ MATCH — OK to import!
HASH SALES 6 CHECK
Script Total: $28.09
Hash Sales 6 Total: $17.09
RESULT: ⚠ MISMATCH — Check before importing!
"""

        self.assertTrue(section_status(log_text, "SALES CHECK"))
        self.assertFalse(section_status(log_text, "HASH SALES"))

    def test_unknown_sales_subdepartment_survives_into_engine_tba_path(self):
        from io import BytesIO

        from openpyxl import load_workbook

        from app import pos_to_quickbooks_v2 as engine

        reports = []
        for report in sms_exports_091426():
            if report.role == "sales":
                rows = (
                    ("Sub-department Single Total",),
                    ("Date:", "09/14/2026", "to", "09/14/2026"),
                    ("", "Sub-Department", "Qty", "Amount", "g/Weight"),
                    (110, "Grocery", 1, "50000.00", "0.00"),
                    (120, "Dairy", 1, "36502.80", "0.00"),
                    (777, "Unmapped Department", 1, "12.34", "0.00"),
                    ("Sales Total", "", "", "86515.14", ""),
                )
                reports.append(replace(report, rows=rows))
            elif report.role == "coupons":
                rows = (
                    ("Sub-department Single Total",),
                    ("Date:", "09/14/2026", "to", "09/14/2026"),
                    ("", "Sub-Department", "Qty", "Amount", "g/Weight"),
                    (110, "Grocery Coupon Distribution", 1, "31.69", "0.00"),
                    (120, "Dairy Coupon Distribution", 1, "30.00", "0.00"),
                    (3542, "Store Coupons", 1, "-1428.18", "0.00"),
                    ("Coupon Distribution Total", "", "", "61.69", ""),
                    ("Store Coupons Total", "", "", "-1428.18", ""),
                )
                reports.append(replace(report, rows=rows))
            else:
                reports.append(report)

        bundle = build_sms_export_bundle(reports)
        data = build_sms_deposit_data(bundle)
        workbook_bytes = build_reporting_workbook(
            Path("assets/SubDept Single Total Report Template.xlsx"),
            bundle,
            data,
        )
        workbook = load_workbook(BytesIO(workbook_bytes), data_only=True)
        report_sheet = workbook["SubDept Sales Report"]
        unknown_rows = [
            row
            for row in report_sheet.iter_rows(values_only=True)
            if row[0] == 777
        ]
        template = load_workbook(
            Path("assets/SubDept Single Total Report Template.xlsx"),
            data_only=False,
        )

        self.assertEqual(len(unknown_rows), 1)
        self.assertEqual(unknown_rows[0][2], 12.34)
        self.assertEqual(unknown_rows[0][6], 12.34)
        self.assertEqual(report_sheet["C22"]._style, template["SubDept Sales Report"]["C22"]._style)

        fixture_path = Path(__file__).parent / f"_unknown_subdept_{uuid4().hex}.xlsx"
        fixture_path.write_bytes(workbook_bytes)
        try:
            parsed = engine.parse_excel_report(fixture_path, date(2026, 9, 14))
        finally:
            fixture_path.unlink()

        self.assertIn(("Dept 777", 12.34), parsed[1])

    def test_engine_binds_hash_control_to_active_report_date(self):
        import openpyxl

        from app import pos_to_quickbooks_v2 as engine

        workbook = openpyxl.Workbook()
        report = workbook.active
        report.title = "SubDept Sales Report"
        report["J3"] = 86502.80
        for sheet_name, control in (
            ("083126 Hash", 99.99),
            ("091426 Hash", 10.89),
        ):
            sheet = workbook.create_sheet(sheet_name)
            sheet["I1"] = "Amount"
            sheet["E2"] = "Total"
            sheet["I2"] = control

        fixture_path = Path(__file__).parent / f"_multi_hash_{uuid4().hex}.xlsx"
        workbook.save(fixture_path)
        try:
            parsed = engine.parse_excel_report(fixture_path, date(2026, 9, 14))
        finally:
            fixture_path.unlink()

        self.assertEqual(parsed[-1], 10.89)

    def test_engine_reads_real_hash_control_total_from_dated_source_tab(self):
        from app import pos_to_quickbooks_v2 as engine

        bundle = build_sms_export_bundle(sms_exports_091426())
        reports = dict(bundle.reports)
        reports["hash"] = replace(
            reports["hash"],
            rows=(
                (("Sub-department Single Total",) + ("",) * 11),
                (("", "", "Sub-Department", "", "", "", "", "Qty", "Amount", "Pkg/Weight") + ("",) * 2),
                (("", 23, "Refunded Discounts", "", "", "", 8, 4.89, "", "") + ("",) * 2),
                (("", 32, "PASS THROUGH DONATIONS", "", "", "", 6, 6.00, "", "") + ("",) * 2),
                (("", "", "", "", "Total", "", 14, "", 10.89, "") + ("",) * 2),
            ),
        )
        real_hash_bundle = replace(bundle, reports=reports)
        data = build_sms_deposit_data(real_hash_bundle)
        workbook_bytes = build_reporting_workbook(
            Path("assets/SubDept Single Total Report Template.xlsx"),
            real_hash_bundle,
            data,
        )
        fixture_path = Path(__file__).parent / f"_hash_control_{uuid4().hex}.xlsx"
        fixture_path.write_bytes(workbook_bytes)
        try:
            parsed = engine.parse_excel_report(fixture_path)
        finally:
            fixture_path.unlink()

        self.assertEqual(parsed[-1], 10.89)

    def test_present_zero_hash_control_matches_zero_activity(self):
        from app import pos_to_quickbooks_v2 as engine

        reports = []
        for report in sms_exports_091426():
            if report.role == "hash":
                rows = (
                    ("Sub-department Single Total",),
                    ("Date:", "09/14/2026", "to", "09/14/2026"),
                    ("", "Sub-Department", "Qty", "Amount", "g/Weight"),
                    (23, "Refunded Discounts", 0, "0.00", "0.00"),
                    (32, "Pass Through Donations", 0, "0.00", "0.00"),
                    (34, "Paid In", 0, "0.00", "0.00"),
                    ("HASH Detail Total", "", "", "0.00", ""),
                )
                reports.append(replace(report, rows=rows))
            else:
                reports.append(report)

        bundle = build_sms_export_bundle(reports)
        data = build_sms_deposit_data(bundle)
        workbook_bytes = build_reporting_workbook(
            Path("assets/SubDept Single Total Report Template.xlsx"),
            bundle,
            data,
        )

        fixture_root = Path(__file__).parent / f"_zero_hash_{uuid4().hex}"
        fixture_root.mkdir()
        workbook_path = fixture_root / "zero-hash.xlsx"
        workbook_path.write_bytes(workbook_bytes)
        old_output_dir = engine.output_dir
        old_log_dir = engine.LOG_DIR
        old_log_disabled = engine.log.disabled
        engine.output_dir = fixture_root
        engine.LOG_DIR = fixture_root
        engine.log.disabled = True
        try:
            parsed = engine.parse_excel_report(workbook_path, bundle.deposit_date)
            hash_activity = engine.parse_hash_sheet(
                workbook_path,
                bundle.deposit_date,
            )
            engine.generate_iif(
                {},
                {},
                {},
                bundle.deposit_date,
                refunded_discounts=hash_activity[0],
                pass_through_total=hash_activity[1],
                paid_in_total=hash_activity[2],
                hash_sales_total=parsed[-1],
            )
            status_text = (fixture_root / "last_run_status.txt").read_text(
                encoding="utf-8"
            )
        finally:
            engine.output_dir = old_output_dir
            engine.LOG_DIR = old_log_dir
            engine.log.disabled = old_log_disabled
            for generated_file in fixture_root.iterdir():
                generated_file.unlink()
            fixture_root.rmdir()

        self.assertEqual(parsed[-1], 0.0)
        self.assertEqual(hash_activity, (0.0, 0.0, 0.0))
        self.assertIn("HASH SALES: ✓ MATCH   $0.00", status_text)
        self.assertIn("✓ ALL CHECKS PASSED", status_text)
        self.assertNotIn("NO EXCEL TOTAL FOUND", status_text)

    def test_signed_hash_activity_matches_its_net_control_total(self):
        from app import pos_to_quickbooks_v2 as engine

        fixture_root = Path(__file__).parent / f"_signed_hash_{uuid4().hex}"
        fixture_root.mkdir()
        old_output_dir = engine.output_dir
        old_log_dir = engine.LOG_DIR
        old_log_disabled = engine.log.disabled
        engine.output_dir = fixture_root
        engine.LOG_DIR = fixture_root
        engine.log.disabled = True
        try:
            engine.generate_iif(
                {},
                {},
                {},
                date(2026, 9, 21),
                refunded_discounts=-22.59,
                pass_through_total=5.50,
                hash_sales_total=-17.09,
            )
            status_text = (fixture_root / "last_run_status.txt").read_text(
                encoding="utf-8"
            )
        finally:
            engine.output_dir = old_output_dir
            engine.LOG_DIR = old_log_dir
            engine.log.disabled = old_log_disabled
            for generated_file in fixture_root.iterdir():
                generated_file.unlink()
            fixture_root.rmdir()

        self.assertIn("HASH SALES: ✓ MATCH   $17.09", status_text)
        self.assertIn("✓ ALL CHECKS PASSED", status_text)

    def test_unmapped_hash_activity_is_preserved_as_tba_and_matches_control(self):
        import openpyxl

        from app import pos_to_quickbooks_v2 as engine

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "092326 Hash"
        sheet.append(["Sub-department Single Total"])
        sheet.append(["Date:", "9/23/2026", "to", "9/23/2026"])
        sheet.append(["S-Dept.", 0, "to", 999999])
        sheet.append(["Tlz.:", 6, "to", 6])
        sheet.append([None, None, "Sub-Department", None, None, None, None, "Qty", "Amount"])
        sheet.append([None, 23, "Refunded Discounts", None, None, None, 2, -16.77])
        sheet.append([None, 32, "PASS THROUGH DONATIONS", None, None, None, 2.4, 2.40])
        sheet.append([None, 40, "Postage", None, None, None, 1, 0.69])
        sheet.append([None, None, None, None, "Total", None, 5.4, None, -13.68])

        fixture_root = Path(__file__).parent / f"_hash_postage_{uuid4().hex}"
        fixture_root.mkdir()
        workbook_path = fixture_root / "hash-postage.xlsx"
        workbook.save(workbook_path)
        old_output_dir = engine.output_dir
        old_log_dir = engine.LOG_DIR
        old_log_disabled = engine.log.disabled
        engine.output_dir = fixture_root
        engine.LOG_DIR = fixture_root
        engine.log.disabled = True
        try:
            details = engine.parse_hash_sheet_details(
                workbook_path,
                date(2026, 9, 23),
            )
            iif_path = engine.generate_iif(
                {},
                {},
                {},
                date(2026, 9, 23),
                refunded_discounts=details["refunded_discounts"],
                pass_through_total=details["pass_through_total"],
                paid_in_total=details["paid_in_total"],
                hash_sales_total=-13.68,
                hash_misc_lines=details["misc_lines"],
            )
            iif_text = iif_path.read_text(encoding="utf-8")
            status_text = (fixture_root / "last_run_status.txt").read_text(
                encoding="utf-8"
            )
        finally:
            engine.output_dir = old_output_dir
            engine.LOG_DIR = old_log_dir
            engine.log.disabled = old_log_disabled
            for generated_file in fixture_root.iterdir():
                generated_file.unlink()
            fixture_root.rmdir()

        self.assertEqual(details["refunded_discounts"], -16.77)
        self.assertEqual(details["pass_through_total"], 2.40)
        self.assertEqual(details["paid_in_total"], 0.0)
        self.assertEqual(details["misc_lines"], [("HASH 40 · Postage", 0.69)])
        self.assertIn("4444 · TBA Purchases", iif_text)
        self.assertIn("HASH 40 · Postage", iif_text)
        self.assertIn("\t-0.69\tHASH 40 · Postage\t", iif_text)
        self.assertIn("HASH SALES: ✓ MATCH   $13.68", status_text)
        self.assertIn("✓ ALL CHECKS PASSED", status_text)

    def test_absent_hash_control_is_none(self):
        import openpyxl

        from app import pos_to_quickbooks_v2 as engine

        workbook = openpyxl.Workbook()
        workbook.active.title = "SubDept Sales Report"
        workbook_path = Path(__file__).parent / f"_no_hash_{uuid4().hex}.xlsx"
        workbook.save(workbook_path)
        workbook.close()
        try:
            parsed = engine.parse_excel_report(workbook_path, date(2026, 9, 14))
        finally:
            workbook_path.unlink()

        self.assertIsNone(parsed[-1])

    def test_generated_workbook_upload_matches_engine_file_interface(self):
        upload_type = streamlit_definition(
            "GeneratedWorkbookUpload", {"dataclass": __import__("dataclasses").dataclass}
        )

        upload = upload_type("daily.xlsx", b"workbook bytes")

        self.assertEqual(upload.name, "daily.xlsx")
        self.assertEqual(upload.getvalue(), b"workbook bytes")

    def test_review_warning_keeps_completed_reporting_workbook_attached(self):
        from types import SimpleNamespace

        complete_sms_run = streamlit_definition(
            "complete_sms_run",
            {
                "reporting_workbook_name": (
                    lambda _deposit_date: "SubDept Single Total Report 9-23-26.xlsx"
                ),
            },
        )

        completed = complete_sms_run(
            {
                "iif_path": Path("deposit_20260923.iif"),
                "iif_bytes": b"balanced iif requiring review",
                "validation": {"all_ok": False, "iif_ok": True},
            },
            SimpleNamespace(deposit_date=date(2026, 9, 23)),
            None,
            reporting_workbook_bytes=b"completed report",
        )

        self.assertEqual(
            completed["reporting_workbook_bytes"], b"completed report"
        )
        self.assertEqual(
            completed["reporting_workbook_name"],
            "SubDept Single Total Report 9-23-26.xlsx",
        )

    def test_partial_sms_roles_are_visible_but_bundle_waits_for_all_six(self):
        reports = {report.filename: report for report in sms_exports_091426()}

        class Uploaded:
            def __init__(self, report):
                self.name = report.filename
                self._data = report.role.encode("ascii")

            def getvalue(self):
                return self._data

        uploads = [Uploaded(report) for report in sms_exports_091426()]

        def parse_export(filename, _data):
            return reports[filename]

        def validate_exports(files):
            return build_sms_export_bundle(
                parse_export(item.name, item.getvalue()) for item in files
            )

        inspect_sms_uploads = streamlit_definition(
            "inspect_sms_uploads",
            {
                "parse_sms_export": parse_export,
                "validate_sms_exports": validate_exports,
            },
        )

        partial = inspect_sms_uploads(uploads[:2])
        complete = inspect_sms_uploads(uploads)

        self.assertEqual(set(partial["reports"]), {"sales", "coupons"})
        self.assertIsNone(partial["bundle"])
        self.assertIsInstance(partial["error"], ValueError)
        self.assertEqual(tuple(complete["reports"]), tuple(complete["bundle"].reports))
        self.assertEqual(
            complete["source_bytes"],
            {report.role: report.role.encode("ascii") for report in sms_exports_091426()},
        )
        self.assertIsNone(complete["error"])

    def test_reopening_manual_step_clears_choice_so_it_cannot_immediately_recomplete(self):
        import app.guided_deposit_state as guided_state

        reopen = getattr(guided_state, "reopen_step_for_edit", None)
        self.assertIsNotNone(reopen)

        required_steps = ("member_shares", "paid_in", "closeout")
        completions = {
            "member_shares": "app",
            "paid_in": "quickbooks",
            "closeout": "app",
        }
        workbook_key = "workbook-123"
        handling_key = f"activity_handling_paid_in_{workbook_key}"
        session_state = {
            handling_key: "Finish manually in QuickBooks",
            f"activity_saved_rows_paid_in_{workbook_key}": [{"amount": 100.0}],
        }

        reopened = reopen(
            session_state,
            required_steps=required_steps,
            completions=completions,
            step="paid_in",
            workbook_key=workbook_key,
        )

        self.assertEqual(reopened, {"member_shares": "app"})
        self.assertNotIn(handling_key, session_state)
        self.assertEqual(
            session_state[f"activity_saved_rows_paid_in_{workbook_key}"],
            [{"amount": 100.0}],
        )

    def test_reopening_app_breakdown_preserves_choice_and_saved_entries(self):
        import app.guided_deposit_state as guided_state

        reopen = getattr(guided_state, "reopen_step_for_edit", None)
        self.assertIsNotNone(reopen)

        workbook_key = "workbook-123"
        handling_key = f"coupon_handling_{workbook_key}"
        session_state = {
            handling_key: "Breakdown in app using Closeout Sheet",
            f"coupon_saved_payload_{workbook_key}": {"mode": "closeout"},
        }

        reopened = reopen(
            session_state,
            required_steps=("coupons", "closeout"),
            completions={"coupons": "app", "closeout": "app"},
            step="coupons",
            workbook_key=workbook_key,
        )

        self.assertEqual(reopened, {})
        self.assertEqual(
            session_state[handling_key],
            "Breakdown in app using Closeout Sheet",
        )
        self.assertEqual(
            session_state[f"coupon_saved_payload_{workbook_key}"],
            {"mode": "closeout"},
        )

    def test_reopening_member_shares_restores_saved_app_choice(self):
        from app.guided_deposit_state import reopen_step_for_edit

        workbook_key = "workbook-123"
        session_state = {
            f"membership_saved_choice_{workbook_key}": (
                "Breakdown in app using the Ownership Payments sheet"
            ),
            f"membership_saved_payments_{workbook_key}": [{"amount": 100.0}],
        }

        reopen_step_for_edit(
            session_state,
            required_steps=("member_shares", "closeout"),
            completions={"member_shares": "app", "closeout": "app"},
            step="member_shares",
            workbook_key=workbook_key,
        )

        self.assertEqual(
            session_state[f"membership_handling_{workbook_key}"],
            "Breakdown in app using the Ownership Payments sheet",
        )
        self.assertEqual(
            session_state[f"membership_saved_payments_{workbook_key}"],
            [{"amount": 100.0}],
        )

    def test_reopening_activity_restores_saved_rows_into_editable_fields(self):
        from app.guided_deposit_state import reopen_step_for_edit

        workbook_key = "workbook-123"
        session_state = {
            f"activity_saved_section_donation_{workbook_key}": {
                "mode": "app",
                "rows": [{
                    "account": "8504000 · Education",
                    "given_to": "Food Pantry",
                    "purpose": "Groceries",
                    "manager": "KC",
                    "amount": 75.0,
                }],
            },
            f"activity_saved_section_paid_out_{workbook_key}": {
                "mode": "app",
                "rows": [{
                    "type": "esp",
                    "original_date": "2026-08-28",
                    "initials": "MR",
                    "amount": 40.0,
                }],
            },
        }

        for step in ("donation", "paid_out"):
            reopen_step_for_edit(
                session_state,
                required_steps=("donation", "paid_out", "closeout"),
                completions={
                    "donation": "app",
                    "paid_out": "app",
                    "closeout": "app",
                },
                step=step,
                workbook_key=workbook_key,
            )

        donation_id = session_state[
            f"activity_row_ids_donation_{workbook_key}"
        ][0]
        self.assertEqual(
            session_state[f"activity_donation_{donation_id}_given_to"],
            "Food Pantry",
        )
        self.assertEqual(
            session_state[f"activity_donation_{donation_id}_account"],
            "8504000 · Education",
        )
        self.assertEqual(
            session_state[f"activity_donation_{donation_id}_amount"],
            75.0,
        )
        paid_out_id = session_state[
            f"activity_row_ids_paid_out_{workbook_key}"
        ][0]
        self.assertEqual(
            session_state[f"activity_paid_out_{paid_out_id}_type"],
            "ESP Deposit",
        )
        self.assertEqual(
            session_state[f"activity_paid_out_{paid_out_id}_date"],
            date(2026, 8, 28),
        )
        self.assertEqual(
            session_state[f"activity_paid_out_{paid_out_id}_account"],
            "1230000 · Miscellaneous Receivable",
        )
        self.assertEqual(
            session_state[f"activity_paid_out_{paid_out_id}_memo"],
            "8/28's ESP Deposit - MR",
        )
        self.assertEqual(
            session_state[f"activity_paid_out_{paid_out_id}_amount"],
            40.0,
        )

    def test_reopening_other_paid_item_restores_selected_account(self):
        from app.guided_deposit_state import reopen_step_for_edit

        workbook_key = "workbook-123"
        session_state = {
            f"activity_saved_section_paid_out_{workbook_key}": {
                "mode": "app",
                "rows": [{
                    "type": "other",
                    "account": "8428000 · Admn - Office Supplies",
                    "memo": "Printer paper",
                    "amount": 40.0,
                }],
            },
        }

        reopen_step_for_edit(
            session_state,
            required_steps=("paid_out", "closeout"),
            completions={"paid_out": "app", "closeout": "app"},
            step="paid_out",
            workbook_key=workbook_key,
        )

        paid_out_id = session_state[
            f"activity_row_ids_paid_out_{workbook_key}"
        ][0]
        self.assertEqual(
            session_state[f"activity_paid_out_{paid_out_id}_type"],
            "Other",
        )
        self.assertEqual(
            session_state[f"activity_paid_out_{paid_out_id}_account"],
            "8428000 · Admn - Office Supplies",
        )
        self.assertEqual(
            session_state[f"activity_paid_out_{paid_out_id}_amount"],
            40.0,
        )

    def test_reopening_coupons_restores_saved_totals_into_editable_fields(self):
        from app.guided_deposit_state import reopen_step_for_edit

        workbook_key = "workbook-123"
        session_state = {
            f"coupon_saved_payload_{workbook_key}": {
                "mode": "closeout",
                "closeout_total": 183.25,
                "ncg_total": 146.0,
                "mfg_total": 37.25,
            }
        }

        reopen_step_for_edit(
            session_state,
            required_steps=("coupons", "closeout"),
            completions={"coupons": "app", "closeout": "app"},
            step="coupons",
            workbook_key=workbook_key,
        )

        self.assertEqual(
            session_state[f"coupon_handling_{workbook_key}"],
            "Breakdown in app using Closeout Sheet",
        )
        self.assertEqual(session_state[f"coupon_closeout_{workbook_key}"], 183.25)
        self.assertEqual(session_state[f"coupon_ncg_{workbook_key}"], 146.0)
        self.assertEqual(session_state[f"coupon_mfg_{workbook_key}"], 37.25)

    def workflow_api(self):
        try:
            from app.deposit_workflow import (
                STEP_CLOSEOUT,
                STEP_INHOUSE,
                active_deposit_step,
                complete_deposit_step,
                deposit_step_rows,
                deposit_workflow_complete,
                edit_deposit_step,
                normalize_step_completions,
                required_deposit_steps,
            )
        except ImportError as exc:
            self.fail(f"guided deposit workflow is missing: {exc}")
        return locals()

    def test_required_steps_use_fixed_gap_free_business_order(self):
        api = self.workflow_api()
        steps = api["required_deposit_steps"](8.45, {"donation": 25, "paid_in": 100, "paid_out": 40}, 18.49, 140.82)
        self.assertEqual(steps, ("member_shares", "donation", "paid_in", "paid_out", "coupons", "inhouse_charges", "closeout"))

    def test_inhouse_follows_coupons_when_only_those_steps_are_required(self):
        api = self.workflow_api()
        self.assertEqual(
            api["required_deposit_steps"](0, {"donation": 0, "paid_in": 0, "paid_out": 0}, 18.49, 140.82),
            ("coupons", "inhouse_charges", "closeout"),
        )

    def test_workflow_detection_reads_charge_house_and_blocks_on_failure(self):
        source = SOURCE.read_text(encoding="utf-8")
        self.assertIn("charge_house_total = read_charge_house_total(", source)
        self.assertIn("workflow_detection_valid = activity_detection_valid and charge_house_valid", source)
        self.assertIn("detection_valid=workflow_detection_valid", source)
        self.assertIn("Could not read Charge (House) from the Balance Sheet", source)

    def test_only_paid_in_still_ends_with_closeout(self):
        api = self.workflow_api()
        self.assertEqual(api["required_deposit_steps"](0, {"donation": 0, "paid_in": 100, "paid_out": 0}, 0, 0), ("paid_in", "closeout"))

    def test_no_optional_activity_still_requires_closeout(self):
        api = self.workflow_api()
        self.assertEqual(api["required_deposit_steps"](0, {"donation": 0, "paid_in": 0, "paid_out": 0}, 0, 0), ("closeout",))

    def test_every_optional_combination_preserves_business_order(self):
        from itertools import product
        api = self.workflow_api()
        ordered = ("member_shares", "donation", "paid_in", "paid_out", "coupons", "inhouse_charges")
        for enabled in product((False, True), repeat=len(ordered)):
            with self.subTest(enabled=enabled):
                steps = api["required_deposit_steps"](1 if enabled[0] else 0, {"donation": 1 if enabled[1] else 0, "paid_in": 1 if enabled[2] else 0, "paid_out": 1 if enabled[3] else 0}, 1 if enabled[4] else 0, 1 if enabled[5] else 0)
                self.assertEqual(steps, tuple(step for step, is_enabled in zip(ordered, enabled) if is_enabled) + ("closeout",))

    def test_completion_methods_advance_to_first_incomplete_step(self):
        api = self.workflow_api(); steps = ("paid_in", "paid_out", "closeout")
        completed = api["complete_deposit_step"](steps, {}, "paid_in", "quickbooks")
        self.assertEqual(completed, {"paid_in": "quickbooks"})
        self.assertEqual(api["active_deposit_step"](steps, completed), "paid_out")
        completed = api["complete_deposit_step"](steps, completed, "paid_out", "app")
        self.assertEqual(api["active_deposit_step"](steps, completed), "closeout")

    def test_normalization_ignores_stale_and_invalid_completion_entries(self):
        api = self.workflow_api()
        normalized = api["normalize_step_completions"](("paid_in", "closeout"), {"donation": "app", "paid_in": "invalid", "closeout": "quickbooks"})
        self.assertEqual(normalized, {"closeout": "quickbooks"})

    def test_edit_preserves_later_activity_but_invalidates_step_and_closeout(self):
        api = self.workflow_api(); steps = ("donation", "paid_in", "paid_out", "closeout")
        completed = {"donation": "app", "paid_in": "app", "paid_out": "quickbooks", "closeout": "app"}
        edited = api["edit_deposit_step"](steps, completed, "paid_in")
        self.assertEqual(edited, {"donation": "app", "paid_out": "quickbooks"})
        self.assertEqual(api["active_deposit_step"](steps, edited), "paid_in")

    def test_editing_closeout_preserves_every_earlier_completion(self):
        api = self.workflow_api()
        steps = ("paid_in", "paid_out", "closeout")
        completed = {
            "paid_in": "app",
            "paid_out": "quickbooks",
            "closeout": "app",
        }
        self.assertEqual(
            api["edit_deposit_step"](steps, completed, "closeout"),
            {"paid_in": "app", "paid_out": "quickbooks"},
        )

    def test_save_inhouse_transition_persists_canonical_payload_and_completion(self):
        from app.guided_deposit_state import save_inhouse_transition

        saved_key = "inhouse_saved_payload_workbook-123"
        transition = save_inhouse_transition(
            ("inhouse_charges", "closeout"),
            {},
            saved_key,
            {
                "actual": 8,
                "rows": [{
                    "account": "8320000 · Store Supplies",
                    "memo": "End of Day - BS",
                    "amount": 8,
                }],
            },
        )
        self.assertEqual(transition["completions"], {"inhouse_charges": "app"})
        self.assertEqual(transition["saved_payload"], {
            saved_key: {
                "actual": 8.0,
                "rows": [{
                    "account": "8320000 · Store Supplies",
                    "memo": "End of Day - BS",
                    "amount": 8.0,
                }],
            }
        })

    def test_saved_inhouse_handoff_locks_actual_and_preserves_over_short(self):
        from app.closeout_reconciliation import (
            STANDARD_CLOSEOUT_ORDER,
            build_closeout_form_payload,
            build_standard_reconciliation,
        )
        from app.guided_step_ui import saved_inhouse_for_closeout

        workbook_key = "workbook-123"
        saved_rows = [{
            "account": "8320000 · Store Supplies",
            "memo": "End of Day - BS",
            "amount": 8,
        }]
        session_state = {f"inhouse_saved_payload_{workbook_key}": {
            "actual": 8, "rows": saved_rows,
        }}
        inhouse = saved_inhouse_for_closeout(
            session_state, workbook_key, ("inhouse_charges", "closeout")
        )
        baselines = {key: 0 for key in STANDARD_CLOSEOUT_ORDER}
        baselines["charge_house"] = 10
        actuals = dict(baselines)
        actuals["charge_house"] = inhouse["actual"]
        payload = build_closeout_form_payload(
            baselines=baselines, actuals=actuals, reviewed=True,
            payroll=0, safe_type="none", safe_amount=0,
            plants_purchase=0, custom_tba=[], final_total=100,
            approve_final_pos=False, inhouse_charges=inhouse["rows"],
        )

        self.assertEqual(payload["actuals"]["charge_house"], 8.0)
        self.assertEqual(payload["inhouse_charges"], saved_rows)
        charge_row = next(row for row in build_standard_reconciliation(
            baselines, payload["actuals"]
        ) if row["key"] == "charge_house")
        self.assertEqual(-charge_row["adjustment_qb_effect"], 2.0)

    def test_inhouse_resave_requires_new_manual_closeout_choice(self):
        from app.guided_step_ui import prepare_closeout_after_inhouse_save

        workbook_key = "book"
        state = {
            f"closeout_payload_{workbook_key}": {"mode": "manual"},
            f"closeout_handling_{workbook_key}": "Finish manually in QuickBooks",
            f"closeout_form_needs_hydration_{workbook_key}": True,
            f"closeout_preview_{workbook_key}": {"stale": True},
        }

        prepare_closeout_after_inhouse_save(state, workbook_key)

        self.assertEqual(state[f"closeout_payload_{workbook_key}"], {"mode": "manual"})
        self.assertNotIn(f"closeout_handling_{workbook_key}", state)
        self.assertNotIn(f"closeout_form_needs_hydration_{workbook_key}", state)
        self.assertNotIn(f"closeout_preview_{workbook_key}", state)

    def test_inhouse_resave_hydrates_prior_app_closeout_without_approval(self):
        from app.guided_step_ui import prepare_closeout_after_inhouse_save

        workbook_key = "book"
        state = {
            f"closeout_payload_{workbook_key}": {"mode": "closeout"},
            f"closeout_handling_{workbook_key}": "Breakdown in app using Closeout Sheet",
            f"closeout_approve_final_{workbook_key}": True,
        }

        prepare_closeout_after_inhouse_save(state, workbook_key)

        self.assertTrue(state[f"closeout_form_needs_hydration_{workbook_key}"])
        self.assertNotIn(f"closeout_handling_{workbook_key}", state)
        self.assertNotIn(f"closeout_approve_final_{workbook_key}", state)

    def test_required_inhouse_handoff_rejects_missing_or_invalid_saved_payload(self):
        from app.guided_step_ui import saved_inhouse_for_closeout

        with self.assertRaisesRegex(ValueError, "InHouse Charges need review"):
            saved_inhouse_for_closeout({}, "book", ("inhouse_charges", "closeout"))
        with self.assertRaisesRegex(ValueError, "InHouse Charges need review"):
            saved_inhouse_for_closeout(
                {"inhouse_saved_payload_book": {"actual": 8, "rows": []}},
                "book", ("inhouse_charges", "closeout"),
            )
        self.assertEqual(saved_inhouse_for_closeout({}, "book", ("closeout",)), {
            "actual": 0.0, "rows": [],
        })

    def test_historical_closeout_seeds_inhouse_widgets_once_without_completion(self):
        from app.closeout_reconciliation import STANDARD_CLOSEOUT_ORDER
        from app.guided_step_ui import seed_historical_inhouse_widgets

        workbook_key = "book"
        actuals = {key: 0 for key in STANDARD_CLOSEOUT_ORDER}
        actuals["charge_house"] = 8
        state = {f"closeout_payload_{workbook_key}": {
            "mode": "closeout", "reviewed": True, "actuals": actuals,
            "payroll": 0, "safe": {"type": "none", "amount": 0},
            "plants_purchase": 0, "custom_tba": [],
            "inhouse_charges": [{
                "account": "8320000 · Store Supplies",
                "memo": "End of Day", "amount": 8,
            }],
            "final_total": 100, "approve_final_pos": False,
        }}
        self.assertTrue(seed_historical_inhouse_widgets(state, workbook_key))
        self.assertEqual(state[f"inhouse_actual_{workbook_key}"], 8.0)
        row_id = state[f"inhouse_ids_{workbook_key}"][0]
        self.assertEqual(state[f"inhouse_account_{workbook_key}_{row_id}"], "8320000 · Store Supplies")
        self.assertEqual(state[f"inhouse_memo_type_{workbook_key}_{row_id}"], "End of Day")
        self.assertEqual(state[f"inhouse_memo_value_{workbook_key}_{row_id}"], "")
        self.assertEqual(state[f"inhouse_amount_{workbook_key}_{row_id}"], 8.0)
        state[f"inhouse_memo_value_{workbook_key}_{row_id}"] = "KC"
        self.assertFalse(seed_historical_inhouse_widgets(state, workbook_key))
        self.assertEqual(state[f"inhouse_memo_value_{workbook_key}_{row_id}"], "KC")
        self.assertNotIn(f"inhouse_saved_payload_{workbook_key}", state)

    def test_reopened_inhouse_restores_actual_and_each_memo_mode(self):
        from app.guided_deposit_state import _hydrate_reopened_app_step, reopen_step_for_edit

        workbook_key = "workbook-123"
        session_state = {
            f"inhouse_saved_payload_{workbook_key}": {
                "actual": 18,
                "rows": [
                    {"account": "8320000 · Store Supplies", "memo": "End of Day - BS", "amount": 8},
                    {"account": "8428000 · Office Supplies", "memo": "Board lunch", "amount": 6},
                    {"account": "8320000 · Store Supplies", "memo": "End of Day", "amount": 4},
                ],
            }
        }
        self.assertTrue(_hydrate_reopened_app_step(session_state, "inhouse_charges", workbook_key))
        row_ids = session_state[f"inhouse_ids_{workbook_key}"]
        self.assertEqual(row_ids, ["restored_0", "restored_1", "restored_2"])
        self.assertEqual(session_state[f"inhouse_actual_{workbook_key}"], 18.0)
        for row_id, account, memo_type, memo_value, amount in (
            ("restored_0", "8320000 · Store Supplies", "End of Day", "BS", 8.0),
            ("restored_1", "8428000 · Office Supplies", "Custom", "Board lunch", 6.0),
            ("restored_2", "8320000 · Store Supplies", "End of Day", "", 4.0),
        ):
            self.assertEqual(session_state[f"inhouse_account_{workbook_key}_{row_id}"], account)
            self.assertEqual(session_state[f"inhouse_memo_type_{workbook_key}_{row_id}"], memo_type)
            self.assertEqual(session_state[f"inhouse_memo_value_{workbook_key}_{row_id}"], memo_value)
            self.assertEqual(session_state[f"inhouse_amount_{workbook_key}_{row_id}"], amount)

        reopened = reopen_step_for_edit(
            session_state,
            required_steps=("inhouse_charges", "closeout"),
            completions={"inhouse_charges": "app", "closeout": "app"},
            step="inhouse_charges",
            workbook_key=workbook_key,
        )
        self.assertEqual(reopened, {})
        self.assertEqual(session_state[f"inhouse_ids_{workbook_key}"], row_ids)

    def test_malformed_saved_inhouse_payload_does_not_restore_or_complete(self):
        from app.guided_deposit_state import _hydrate_reopened_app_step, reopen_step_for_edit

        workbook_key = "workbook-123"
        session_state = {
            f"inhouse_saved_payload_{workbook_key}": {
                "actual": 8,
                "rows": [{"account": "8320000 · Store Supplies", "memo": "Board lunch", "amount": 7}],
            }
        }
        self.assertFalse(_hydrate_reopened_app_step(session_state, "inhouse_charges", workbook_key))
        self.assertNotIn(f"inhouse_actual_{workbook_key}", session_state)
        self.assertEqual(
            reopen_step_for_edit(
                session_state,
                required_steps=("inhouse_charges", "closeout"),
                completions={"inhouse_charges": "app", "closeout": "app"},
                step="inhouse_charges",
                workbook_key=workbook_key,
            ),
            {},
        )

    def test_editing_inhouse_requires_closeout_review_again(self):
        api = self.workflow_api()
        self.assertEqual(
            api["edit_deposit_step"](
                ("coupons", "inhouse_charges", "closeout"),
                {"coupons": "app", "inhouse_charges": "app", "closeout": "app"},
                "inhouse_charges",
            ),
            {"coupons": "app"},
        )

    def test_reopened_closeout_hydrates_canonical_form_and_drops_preview(self):
        from app.closeout_reconciliation import STANDARD_CLOSEOUT_ORDER
        from app.guided_deposit_state import reopen_step_for_edit

        workbook_key = "workbook-123"
        payload_key = f"closeout_payload_{workbook_key}"
        preview_key = f"closeout_preview_{workbook_key}"
        session_state = {
            payload_key: {
                "mode": "closeout",
                "reviewed": True,
                "actuals": {
                    field: float(index + 10)
                    for index, field in enumerate(STANDARD_CLOSEOUT_ORDER)
                },
                "payroll": -4000.0,
                "safe": {"type": "overage", "amount": 7.5},
                "plants_purchase": 12.34,
                "custom_tba": [
                    {
                        "memo": "Printer repair",
                        "amount": 45.67,
                        "direction": "adds",
                    }
                ],
                "inhouse_charges": [
                    {
                        "account": "8320000 · Store Supplies",
                        "memo": "End of Day",
                        "amount": 13.0,
                    }
                ],
                "final_total": 1000.0,
                "approve_final_pos": True,
            },
            preview_key: {"input_fingerprint": "stale", "preview": {}},
        }

        reopened = reopen_step_for_edit(
            session_state,
            required_steps=("closeout",),
            completions={"closeout": "app"},
            step="closeout",
            workbook_key=workbook_key,
        )

        self.assertEqual(reopened, {})
        self.assertNotIn(preview_key, session_state)
        self.assertEqual(
            session_state[f"closeout_handling_{workbook_key}"],
            "Breakdown in app using Closeout Sheet",
        )
        self.assertTrue(session_state[f"closeout_reviewed_{workbook_key}"])
        self.assertEqual(
            session_state[f"closeout_actual_cash_{workbook_key}"],
            10.0,
        )
        self.assertEqual(
            session_state[f"closeout_payroll_{workbook_key}"],
            "Removes $4,000",
        )
        self.assertEqual(
            session_state[f"closeout_safe_type_{workbook_key}"], "Overage")
        self.assertEqual(
            session_state[f"closeout_safe_amount_{workbook_key}"], 7.5
        )
        self.assertEqual(
            session_state[f"closeout_plants_{workbook_key}"], 12.34)
        self.assertEqual(session_state[f"closeout_custom_ids_{workbook_key}"], [0])
        self.assertEqual(session_state[f"closeout_inhouse_ids_{workbook_key}"], [0])
        self.assertEqual(
            session_state[f"closeout_inhouse_account_{workbook_key}_0"],
            "8320000 · Store Supplies",
        )
        self.assertEqual(
            session_state[f"closeout_inhouse_memo_{workbook_key}_0"],
            "End of Day",
        )
        self.assertEqual(
            session_state[f"closeout_inhouse_amount_{workbook_key}_0"],
            13.0,
        )
        self.assertEqual(
            session_state[f"closeout_custom_memo_{workbook_key}_0"],
            "Printer repair",
        )
        self.assertEqual(
            session_state[f"closeout_custom_direction_{workbook_key}_0"],
            "Adds to deposit",
        )
        self.assertEqual(
            session_state[f"closeout_final_total_{workbook_key}"], 1000.0)
        self.assertNotIn(
            f"closeout_approve_final_{workbook_key}",
            session_state,
        )

    def test_rows_are_gap_free_and_expose_user_facing_statuses(self):
        api = self.workflow_api()
        rows = api["deposit_step_rows"](("paid_in", "closeout"), {"paid_in": "quickbooks"})
        self.assertEqual(rows, [{"number": 1, "step": "paid_in", "label": "Paid In", "status": "Finish in QuickBooks", "complete": True, "current": False}, {"number": 2, "step": "closeout", "label": "Closeout Sheet", "status": "Current", "complete": False, "current": True}])

    def test_final_eligibility_requires_every_step(self):
        api = self.workflow_api(); steps = ("paid_in", "closeout")
        self.assertFalse(api["deposit_workflow_complete"](steps, {"paid_in": "app"}))
        self.assertTrue(api["deposit_workflow_complete"](steps, {"paid_in": "app", "closeout": "quickbooks"}))
