import unittest
from datetime import date
from io import BytesIO
from numbers import Number
from pathlib import Path

from openpyxl import load_workbook

from app.daily_reporting_workbook import (
    build_reporting_workbook,
    reporting_workbook_name,
)
from app.sms_deposit_data import build_sms_deposit_data
from app.sms_exports import SmsExport, build_sms_export_bundle
from tests.sms_fixture_factory import sms_exports_091426


TEMPLATE = Path("assets/SubDept Single Total Report Template.xlsx")
EXPECTED_SHEETS = [
    "Membership Payment Plans",
    "InHouse Purchases",
    "Random Sales",
    "SMS Look Up",
    "SubDept Single",
    "SubDept Coupon (Local Discount)",
    "Maria College 1 Sales Data",
    "Maria College 2 Sales Data",
    "SubDept Sales Report",
    "XXXXXX Discounts",
    "XXXXXX BS",
    "XXXXXX Hash",
]
MONEY_FORMAT = "0.00"


class DailyReportingWorkbookTests(unittest.TestCase):
    def test_report_name_uses_no_leading_zero_month_or_day(self):
        self.assertEqual(
            reporting_workbook_name(date(2026, 9, 4)),
            "SubDept Single Total Report 9-4-26.xlsx",
        )

    def test_template_is_sanitized_without_removing_reference_layout(self):
        workbook = load_workbook(TEMPLATE, data_only=False)

        self.assertEqual(workbook.sheetnames, EXPECTED_SHEETS)
        for sheet_name in (
            "SubDept Single",
            "SubDept Coupon (Local Discount)",
            "XXXXXX Discounts",
            "XXXXXX BS",
            "XXXXXX Hash",
        ):
            sheet = workbook[sheet_name]
            self.assertFalse(
                any(cell.value is not None for row in sheet.iter_rows() for cell in row),
                sheet_name,
            )

        report = workbook["SubDept Sales Report"]
        self.assertEqual(report["A1"].value, "Sub-Dept. Single Total Report")
        self.assertEqual(report["A4"].value, "Sub-Dept. Number")
        self.assertEqual(report["B54"].value, "Totals")
        for row in report.iter_rows(min_row=5, max_row=54, min_col=3, max_col=7):
            self.assertTrue(all(cell.value is None for cell in row))
        for address in ("J1", "J2", "J3", "M1", "M2"):
            self.assertIsNone(report[address].value)
        self.assertFalse(
            any(cell.data_type == "f" for row in report.iter_rows() for cell in row)
        )
        self.assertEqual(report.column_dimensions["B"].width, 26.140625)
        self.assertEqual(report.row_dimensions[1].height, 18.75)
        self.assertEqual(report["A1"].fill.fgColor.rgb, "FFFFFF00")

    def test_generated_091426_report_contains_verified_values_and_dated_tabs(self):
        bundle = build_sms_export_bundle(sms_exports_091426())
        input_dates = tuple(report.report_date for report in bundle.reports.values())
        data = build_sms_deposit_data(bundle)

        output = build_reporting_workbook(TEMPLATE, bundle, data)
        values_workbook = load_workbook(BytesIO(output), data_only=True)
        formula_workbook = load_workbook(BytesIO(output), data_only=False)

        self.assertEqual(
            values_workbook.sheetnames[-3:],
            ["091426 Discounts", "091426 BS", "091426 Hash"],
        )
        expected = {"J3": 86502.80, "M1": 74.00, "M2": -1354.18}
        for address, amount in expected.items():
            with self.subTest(address=address):
                self.assertEqual(
                    values_workbook["SubDept Sales Report"][address].value, amount
                )
                self.assertEqual(
                    formula_workbook["SubDept Sales Report"][address].value, amount
                )

        sales = values_workbook["SubDept Single"]
        self.assertEqual(
            list(sales.iter_rows(min_row=1, max_row=2, max_col=7, values_only=True)),
            [
                (10, "Grocery", None, None, None, 1, 50000.00),
                (20, "Wellness", None, None, None, 1, 36502.80),
            ],
        )
        coupons = values_workbook["SubDept Coupon (Local Discount)"]
        self.assertEqual(
            list(coupons.iter_rows(min_row=1, max_row=3, max_col=7, values_only=True)),
            [
                (10, "Grocery Coupon Distribution", None, None, None, 1, 31.69),
                (20, "Wellness Coupon Distribution", None, None, None, 1, 30.00),
                (3542, "Store Coupons", None, None, None, 1, -1428.18),
            ],
        )

        self.assertEqual(
            values_workbook["091426 Discounts"]["A2"].value, "Date:"
        )
        self.assertEqual(values_workbook["091426 BS"]["A3"].value, "Code")
        self.assertEqual(
            values_workbook["091426 Hash"]["B5"].value, "Pass Through Donations"
        )
        for sheet_name, column, label in (
            ("091426 Discounts", "B", "Member Discount"),
            ("091426 BS", "B", "Milk Bottle Return"),
            ("091426 Hash", "B", "Pass Through Donations"),
        ):
            with self.subTest(sheet=sheet_name, column=column):
                self.assertGreaterEqual(
                    values_workbook[sheet_name].column_dimensions[column].width,
                    len(label) + 2,
                )
        self.assertEqual(
            tuple(report.report_date for report in bundle.reports.values()), input_dates
        )

        template = load_workbook(TEMPLATE, data_only=False)
        for address in ("A1", "A4", "C22", "G54", "J3", "M1", "M2"):
            with self.subTest(style=address):
                self.assertEqual(
                    formula_workbook["SubDept Sales Report"][address]._style,
                    template["SubDept Sales Report"][address]._style,
                )

        for sheet_name, addresses in {
            "SubDept Single": ("G1", "G2"),
            "SubDept Coupon (Local Discount)": ("G1", "G2", "G3"),
            "SubDept Sales Report": ("C22", "G22", "C54", "G54", "J3", "M1", "M2"),
        }.items():
            for address in addresses:
                with self.subTest(sheet=sheet_name, money=address):
                    cell = values_workbook[sheet_name][address]
                    self.assertIsInstance(cell.value, Number)
                    self.assertIn(MONEY_FORMAT, cell.number_format)

        self.assertFalse(
            any(
                cell.data_type == "f"
                for sheet in formula_workbook.worksheets
                for row in sheet.iter_rows()
                for cell in row
            )
        )

    def test_department_snapshot_uses_sales_plus_coupon_values(self):
        sales_rows = (
            ("Sub-department Single Total",),
            ("Date:", "09/14/2026", "to", "09/14/2026"),
            ("", "Sub-Department", "Qty", "Amount", "g/Weight"),
            (110, "Grocery Non-Tax", 1, "50000.00", "0.00"),
            (120, "Refrigerated", 1, "36502.80", "0.00"),
            ("Sales Total", "", "", "86502.80", ""),
        )
        coupon_rows = (
            ("Sub-department Single Total",),
            ("Date:", "09/14/2026", "to", "09/14/2026"),
            ("", "Sub-Department", "Qty", "Amount", "g/Weight"),
            (110, "Grocery Coupon Distribution", 1, "31.69", "0.00"),
            (120, "Refrigerated Coupon Distribution", 1, "30.00", "0.00"),
            (3542, "Store Coupons", 1, "-1428.18", "0.00"),
            ("Coupon Distribution Total", "", "", "61.69", ""),
            ("Store Coupons Total", "", "", "-1428.18", ""),
        )
        reports = list(sms_exports_091426())
        reports = [
            SmsExport(
                role=report.role,
                filename=report.filename,
                report_date=report.report_date,
                rows=sales_rows
                if report.role == "sales"
                else coupon_rows
                if report.role == "coupons"
                else report.rows,
            )
            for report in reports
        ]
        bundle = build_sms_export_bundle(reports)
        data = build_sms_deposit_data(bundle)

        workbook = load_workbook(
            BytesIO(build_reporting_workbook(TEMPLATE, bundle, data)), data_only=False
        )
        report = workbook["SubDept Sales Report"]

        self.assertEqual(report["C22"].value, 50000.00)
        self.assertEqual(report["D22"].value, 31.69)
        self.assertEqual(report["G22"].value, 50031.69)
        self.assertEqual(report["C23"].value, 36502.80)
        self.assertEqual(report["D23"].value, 30.00)
        self.assertEqual(report["G23"].value, 36532.80)


if __name__ == "__main__":
    unittest.main()
